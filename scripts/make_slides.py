"""Generate the BTP defence slide deck.

    python -m scripts.make_slides

Writes docs/BTP_Defence_Slides.pptx — 20 slides, 16:9, following the
Problem → Solution → Architecture → Demo → Results → Conclusion arc.

Like the thesis report, the results slides read from results/ so the deck cannot
quote a number the repository cannot produce.

Deck discipline: one idea per slide, short lines, no paragraphs. The speaker
carries the argument; the slide carries the claim. Speaker notes hold the detail
so the slide does not have to.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "BTP_Defence_Slides.pptx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
NAVY_MID = RGBColor(0x2E, 0x5C, 0x8A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xB2, 0x6B, 0x00)

W, H = Inches(13.333), Inches(7.5)


def load_summaries() -> dict[str, dict]:
    out: dict[str, tuple[str, dict]] = {}
    results = ROOT / "results"
    if not results.exists():
        return {}
    for path in sorted(results.glob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            continue
        if not payload.get("summary") or not payload.get("system"):
            continue
        system, run_at = payload["system"], payload.get("run_at", "")
        if system not in out or run_at > out[system][0]:
            out[system] = (run_at, payload["summary"])
    return {k: v[1] for k, v in out.items()}


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]

    # -- primitives --
    def _slide(self):
        return self.prs.slides.add_slide(self.blank)

    def _box(self, slide, left, top, width, height):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        return tf

    def _bar(self, slide) -> None:
        """Thin navy rule under the title — the deck's only decoration."""
        from pptx.enum.shapes import MSO_SHAPE

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.28), Inches(11.9), Emu(22860)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.fill.background()
        shape.shadow.inherit = False

    def notes(self, slide, text: str) -> None:
        slide.notes_slide.notes_text_frame.text = text

    # -- slide kinds --
    def title_slide(self, title: str, subtitle: str, meta: str):
        slide = self._slide()
        tf = self._box(slide, Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.4))
        p = tf.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb = Pt(40), True, NAVY

        p = tf.add_paragraph()
        p.text = subtitle
        p.font.size, p.font.color.rgb, p.font.italic = Pt(20), NAVY_MID, True
        p.space_before = Pt(16)

        p = tf.add_paragraph()
        p.text = meta
        p.font.size, p.font.color.rgb = Pt(13), MUTED
        p.space_before = Pt(24)
        return slide

    def section(self, label: str, title: str):
        slide = self._slide()
        tf = self._box(slide, Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.8))
        p = tf.paragraphs[0]
        p.text = label.upper()
        p.font.size, p.font.bold, p.font.color.rgb = Pt(14), True, ACCENT
        p = tf.add_paragraph()
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb = Pt(36), True, NAVY
        return slide

    def bullets(self, title: str, items: list, *, kicker: str = "", size: int = 18):
        """items: str, or (str, indent_level), or ('**', str) for an emphasised line."""
        slide = self._slide()
        tf = self._box(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.9))
        p = tf.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, NAVY
        self._bar(slide)

        top = Inches(1.6)
        if kicker:
            ktf = self._box(slide, Inches(0.7), top, Inches(11.9), Inches(0.6))
            kp = ktf.paragraphs[0]
            kp.text = kicker
            kp.font.size, kp.font.color.rgb, kp.font.italic = Pt(16), MUTED, True
            top = Inches(2.25)

        tf = self._box(slide, Inches(0.7), top, Inches(11.9), Inches(4.6))
        first = True
        for item in items:
            emphasis = False
            level = 0
            if isinstance(item, tuple):
                if item[0] == "**":
                    emphasis, text = True, item[1]
                else:
                    text, level = item[0], item[1]
            else:
                text = item
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = ("" if emphasis else "• ") + text
            p.level = level
            p.font.size = Pt(size + 2 if emphasis else (size - 2 if level else size))
            p.font.bold = emphasis
            p.font.color.rgb = ACCENT if emphasis else (MUTED if level else INK)
            p.space_after = Pt(12 if not level else 6)
        return slide

    def quote(self, title: str, quote_text: str, attribution: str = ""):
        slide = self._slide()
        tf = self._box(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.9))
        p = tf.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, NAVY
        self._bar(slide)

        tf = self._box(slide, Inches(1.2), Inches(2.4), Inches(11.0), Inches(3.0))
        p = tf.paragraphs[0]
        p.text = quote_text
        p.font.size, p.font.bold, p.font.color.rgb = Pt(24), True, NAVY_MID
        p.alignment = PP_ALIGN.CENTER
        if attribution:
            p = tf.add_paragraph()
            p.text = attribution
            p.font.size, p.font.color.rgb = Pt(15), MUTED
            p.alignment = PP_ALIGN.CENTER
            p.space_before = Pt(24)
        return slide

    def mono(self, title: str, text: str, *, size: int = 12):
        slide = self._slide()
        tf = self._box(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.9))
        p = tf.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, NAVY
        self._bar(slide)

        tf = self._box(slide, Inches(0.7), Inches(1.65), Inches(11.9), Inches(5.2))
        first = True
        for line in text.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = line
            p.font.name, p.font.size, p.font.color.rgb = "Consolas", Pt(size), INK
            p.space_after = Pt(0)
        return slide

    def table(self, title: str, headers: list[str], rows: list[list[str]],
              *, widths: list[float] | None = None):
        slide = self._slide()
        tf = self._box(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.9))
        p = tf.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, NAVY
        self._bar(slide)

        n_rows, n_cols = len(rows) + 1, len(headers)
        height = Inches(min(5.0, 0.42 * n_rows + 0.2))
        shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.7), Inches(1.7),
                                       Inches(11.9), height)
        table = shape.table
        if widths:
            for i, w in enumerate(widths):
                table.columns[i].width = Inches(w)

        for i, head in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = head
            para = cell.text_frame.paragraphs[0]
            para.font.size, para.font.bold, para.font.color.rgb = Pt(13), True, WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY

        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(value)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(12)
                para.font.color.rgb = INK
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF7, 0xF9, 0xFC) if r % 2 else WHITE
        return slide

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return path


def build() -> Deck:
    d = Deck()
    s = load_summaries()

    def num(system: str, key: str, suffix: str = "") -> str:
        v = (s.get(system) or {}).get(key)
        return "—" if v is None else f"{v:,.1f}{suffix}"

    # 1
    slide = d.title_slide(
        "Does Adversarial Multi-Agent Debate\nImprove LLM Financial Forecasting?",
        "A leakage-free evaluation on Indian equities (NSE)",
        f"B.Tech Capstone Project  ·  {date.today():%B %Y}",
    )
    d.notes(slide, "Open with the question, not the system. The system is the "
                   "instrument; the question is the project.")

    # 2
    slide = d.bullets(
        "The problem",
        [
            "Forecasting needs four incompatible kinds of evidence at once:",
            ("price action · company fundamentals · news sentiment · macro context", 1),
            "One LLM handed all of it:",
            ("hallucinates numbers · cannot fetch market data · collapses conflict "
             "into bland consensus", 1),
            ("**", "The industry answer: split the work across specialist agents, "
             "then make them argue."),
        ],
    )
    d.notes(slide, "Keep this to 40 seconds. Everyone already believes the problem.")

    # 3
    slide = d.bullets(
        "First, the honest part",
        [
            "That architecture is already published.",
            ("TradingAgents — arXiv:2412.20138, Tauric Research, open source", 1),
            ("Specialist analysts · Bull and Bear debate · risk team · synthesiser", 1),
            ("**", "So 'we built a Bull/Bear debate' is not a contribution."),
            "A supervisor who searches the phrase finds that paper in one query.",
            "We say it first — and then say what IS ours.",
        ],
    )
    d.notes(slide, "Say this slide before anyone can ask it. It converts the single "
                   "biggest attack on the project into evidence of rigour.")

    # 4
    slide = d.quote(
        "The research question",
        "Under an evaluation protocol that provably excludes lookahead bias,\n"
        "does adversarial debate beat a single LLM and a no-debate ensemble —\n"
        "and is any gain worth its token cost?",
        "Worth answering either way. A measured 'no' is a real finding.",
    )
    d.notes(slide, "The contribution moves from architecture to MEASUREMENT. "
                   "If they remember one slide, make it this one.")

    # 5
    slide = d.bullets(
        "Problem 1 — the backtest that measures memory",
        [
            "'Test a forecast made three months ago against what happened.'",
            ("The model's weights already encode what happened.", 1),
            "arXiv:2512.23847 — Lookahead Propensity:",
            ("materially positive across the training window", 1),
            ("collapses to ~zero right after the training cutoff", 1),
            ("**", "Only a post-cutoff test window is valid. Everything else is recall."),
        ],
    )
    d.notes(slide, "This is why the test window is a design constraint, not a detail.")

    # 6
    slide = d.bullets(
        "Problem 2 — the phase order that guarantees failure",
        [
            "Build for 10 weeks, evaluate in week 11:",
            ("if the protocol is broken, there is no time left to fix it", 1),
            ("'multi-agent beats one LLM' is unprovable without the baseline", 1),
            ("**", "Rule adopted: build the ruler before the thing you measure."),
            "Evaluation moved to Phase 2. Baseline shipped in Phase 1.",
        ],
    )
    d.notes(slide, "Examiners like hearing that a plan was changed for a reason.")

    # 7
    slide = d.bullets(
        "Why Indian equities only — a choice, not a shortcut",
        [
            "It IS the novelty — the literature is overwhelmingly US large-cap.",
            "One market = one calendar, one currency, one regulator, one filing deadline.",
            ("That is what makes the as_of protocol PROVABLE rather than approximate.", 1),
            "It buys India-specific engineering worth real marks:",
            ("point-in-time bhavcopy universe · SEBI filing lag · NIFTY-relative "
             "scoring · lakh/crore checks", 1),
            "NSE publishes a free, keyless, complete daily archive. Few markets do.",
        ],
    )
    d.notes(slide, "Never call it a limitation. Four reasons, stated as a decision.")

    # 8
    slide = d.mono(
        "Architecture",
        "   Ticker + as-of date\n"
        "          |          <-- India boundary: resolve to .NS/.BO, or REJECT\n"
        "          v\n"
        "  +-------- SPECIALIST LAYER  (parallel fan-out) ---------+\n"
        "  |  Technical   Fundamental   Sentiment   Macro          |\n"
        "  +--------------------------------------------------------+\n"
        "          |  fan-in\n"
        "          v\n"
        "   RECONCILIATION GATE    every number re-checked vs source\n"
        "          |\n"
        "          v\n"
        "   DEBATE STAGE   <---- ABLATABLE (one conditional edge)\n"
        "     Bull -> Bear rebuts -> Bull counters -> Bear closes\n"
        "          |\n"
        "          v\n"
        "   CHIEF RISK OFFICER     weighs evidence, issues the call\n"
        "          |\n"
        "          v\n"
        "   signal | target range (INR) | confidence % | citations\n"
        "          |\n"
        "          v\n"
        "   EVALUATION HARNESS     leakage-free scoring",
        size=13,
    )
    d.notes(slide, "Point at the ablatable edge. That edge is the experiment.")

    # 9
    slide = d.table(
        "Seven agents, seven evidence slices",
        ["Agent", "Sees", "Produces"],
        [
            ["Technical", "OHLCV + indicators", "score −10..+10, levels"],
            ["Fundamental", "SEBI-filtered fundamentals", "valuation, health"],
            ["Sentiment", "Deduped Indian headlines", "polarity, catalysts"],
            ["Macro", "RBI, USD/INR, crude, sectors", "Tailwind/Neutral/Headwind"],
            ["Bull / Bear", "All 4 reports + transcript", "case, target, rebuttal"],
            ["Risk Officer", "All 4 + full debate", "FinalForecast"],
        ],
        widths=[2.2, 4.6, 5.1],
    )
    d.notes(slide, "Each specialist sees ONLY its slice. If all four saw everything "
                   "they'd converge and the ablation would measure nothing.")

    # 10
    slide = d.bullets(
        "Leakage control 1 — as_of is enforced in code",
        [
            "Every data function takes an as_of date and cannot return a later row.",
            "Enforced in the DATA LAYER — never by prompting.",
            ("Prompting is a secondary measure. It is not a control and we don't "
             "claim it is.", 1),
            "Realised outcomes live in ONE quarantined module.",
            ("A test parses every forecasting module's imports and asserts none can "
             "reach it.", 1),
        ],
    )
    d.notes(slide, "The import-graph test is the crispest single answer to 'how do "
                   "you know it doesn't leak?'")

    # 11
    slide = d.bullets(
        "Leakage control 2 — a point-in-time universe",
        [
            "Test universe = stocks that ACTUALLY TRADED on each as_of date.",
            ("read from that day's NSE bhavcopy, not today's NIFTY 50 list", 1),
            "Today's index membership is textbook survivorship bias:",
            ("it silently excludes the delisted, merged and demoted — the losers", 1),
            ("**", "The bhavcopy is a complete daily snapshot, so we can PROVE the "
             "universe was right."),
            "The frozen test set records its exclusions, e.g. TATAMOTORS on 2025-11-06.",
        ],
    )
    d.notes(slide, "Very few student backtests can prove their universe. Say so.")

    # 12
    slide = d.bullets(
        "Leakage control 3 — the SEBI filing rule",
        [
            "SEBI LODR Reg. 33: quarterly results filed within 45 days of quarter end.",
            ("Q3 numbers (quarter ends 31 Dec) were NOT public until ~14 Feb.", 1),
            "Rule: a quarter is visible only if as_of ≥ quarter_end + 45 days.",
            ("**", "And a stricter one the data forced on us:"),
            "yfinance serves TODAY's P/E and market cap — not what was knowable then.",
            ("So in backtest mode those fields are WITHHELD entirely.", 1),
            ("Fundamental evidence is often empty. That is the honest answer, and "
             "we report it.", 1),
        ],
    )
    d.notes(slide, "The willingness to ship an empty evidence block rather than a "
                   "convenient number is the strongest integrity signal in the project.")

    # 13
    slide = d.bullets(
        "The reconciliation gate",
        [
            "Every numeric claim is re-checked against its cached source before it "
            "propagates.",
            ("A hallucinated adjective is embarrassing. A hallucinated NUMBER is "
             "invisible.", 1),
            "Catches: value mismatch · MAGNITUDE (lakh/crore 100×) · inverted ranges "
            "· fabricated citations",
            ("**", "It FLAGS — it never silently repairs."),
            "Clamping a bad target would measure our clamp, not the model.",
            "Unperformable checks are recorded as 'unverifiable', never as passes.",
            "Proven by attack: the test suite injects fakes and asserts each is caught.",
        ],
        size=17,
    )
    d.notes(slide, "If asked for one piece of engineering to defend, defend this one.")

    # 14
    slide = d.bullets(
        "How we score",
        [
            "Directional accuracy — the headline",
            "NIFTY-relative accuracy — reported ALONGSIDE it, always",
            "MAE / MAPE on the target, in rupees",
            "Brier score + reliability curve — calibration",
            "Cost per forecast: tokens, calls, seconds",
            ("**", "Horizon = 21 TRADING SESSIONS, not 21 calendar days."),
            "≥3 repeats per configuration, unique cache nonce each.",
            ("Without the nonce, repeats hit the cache and report zero variance.", 1),
        ],
        size=17,
    )
    d.notes(slide, "The nonce detail always lands well — it shows we thought about "
                   "how our own tooling could fake a result.")

    # 15 — the real measured result
    slide = d.bullets(
        "Why raw accuracy alone would flatter us",
        [
            f"The 'always up' rule — predicts up every single time:",
            (f"raw directional accuracy: {num('naive-alwaysup', 'directional_accuracy_pct', '%')}", 1),
            (f"versus NIFTY 50: {num('naive-alwaysup', 'nifty_relative_accuracy_pct', '%')}  "
             f"— near chance", 1),
            ("**", "A model that says 'Buy' every time can look like a forecaster."),
            "Excess-return scoring removes the illusion.",
            "This is measured, committed, and in the results files today.",
        ],
        kicker="Measured on the frozen test set — no LLM required",
    )
    d.notes(slide, "Real numbers from our own committed results. An examiner WILL "
                   "ask about index drift; this pre-empts it with evidence.")

    # 16
    slide = d.section("Demo", "The system, running")
    d.notes(slide,
            "DEMO ORDER:\n"
            "1. forecast AAPL --as-of ...  -> rejected at the boundary\n"
            "2. --data-only                -> evidence pack, no model calls\n"
            "3. Streamlit: run the full pipeline, show the live agent log\n"
            "4. Bull vs Bear panel, then the forecast card\n"
            "5. Scroll to the reconciliation gate\n\n"
            "EVERY run must be served from the permanent LLM cache. Never let a "
            "live API call decide whether the defence succeeds.")

    # 17
    have_llm = "baseline" in s or "multiagent" in s
    rows = [
        ["Naive — momentum", num("naive-momentum", "directional_accuracy_pct", "%"),
         num("naive-momentum", "nifty_relative_accuracy_pct", "%"),
         num("naive-momentum", "mean_tokens_per_forecast")],
        ["Naive — always up", num("naive-alwaysup", "directional_accuracy_pct", "%"),
         num("naive-alwaysup", "nifty_relative_accuracy_pct", "%"),
         num("naive-alwaysup", "mean_tokens_per_forecast")],
        ["Single-LLM baseline", num("baseline", "directional_accuracy_pct", "%"),
         num("baseline", "nifty_relative_accuracy_pct", "%"),
         num("baseline", "mean_tokens_per_forecast")],
        ["Multi-agent + debate", num("multiagent", "directional_accuracy_pct", "%"),
         num("multiagent", "nifty_relative_accuracy_pct", "%"),
         num("multiagent", "mean_tokens_per_forecast")],
        ["Multi-agent, no debate", num("multiagent-nodebate", "directional_accuracy_pct", "%"),
         num("multiagent-nodebate", "nifty_relative_accuracy_pct", "%"),
         num("multiagent-nodebate", "mean_tokens_per_forecast")],
    ]
    slide = d.table("Results", ["System", "Dir. acc", "vs NIFTY", "Tokens/fc"], rows,
                    widths=[4.6, 2.4, 2.4, 2.5])
    d.notes(slide, "Dashes are configurations not yet run. Regenerate this deck after "
                   "`python -m evaluate ablate --repeats 3` and they fill in."
                   if not have_llm else
                   "Read the deltas against the repeat-to-repeat standard deviation "
                   "before claiming any effect.")

    # 18
    slide = d.bullets(
        "Does the debate earn its cost?",
        [
            "The debate is ONE conditional edge in the graph.",
            ("full system: 9 LLM calls   ·   no-debate: 5 LLM calls", 1),
            ("verified end-to-end against a stubbed model, so the count is exact", 1),
            "We report the accuracy delta and the token ratio side by side.",
            ("**", "A gain of 2 points for 2.4× the cost is a different finding "
             "from a free one."),
            "And we read both against the repeat-to-repeat standard deviation "
            "before calling anything an effect.",
        ],
    )
    d.notes(slide, "This is the answer to 'so what?'. Cost is a first-class result, "
                   "not an afterthought.")

    # 19
    slide = d.bullets(
        "What we do NOT claim",
        [
            "Restatement leakage remains — yfinance serves restated financials.",
            ("The 45-day rule fixes TIMING leakage, not restatement leakage.", 1),
            "India CPI is unavailable free — documented, never estimated.",
            "GDELT gives headline metadata and tone, not article text.",
            "The test set is small — differences of a few points sit inside noise.",
            ("**", "Every one of these is in the report. None was discovered by "
             "an examiner."),
        ],
    )
    d.notes(slide, "Volunteering limitations before they're asked is the single "
                   "cheapest way to gain credibility in a viva.")

    # 20
    slide = d.bullets(
        "Conclusion",
        [
            "Built: a 7-agent forecasting system for NSE equities, end to end.",
            "Built first: the evaluation protocol that makes its results mean anything.",
            ("as_of in the data layer · quarantined outcomes · point-in-time universe "
             "· SEBI lag · reconciliation gate", 1),
            "Every ablation is a configuration, not a code path.",
            "Every number in the report traces to a committed results file.",
            ("**", "The contribution is not the debate. It is knowing whether the "
             "debate works."),
        ],
    )
    d.notes(slide, "End on the reframing you opened with. Then stop talking.")

    return d


def main() -> int:
    deck = build()
    path = deck.save(OUT)
    print(f"Wrote {path}  ({path.stat().st_size / 1024:.0f} KB, "
          f"{len(deck.prs.slides.__iter__.__self__._sldIdLst)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
