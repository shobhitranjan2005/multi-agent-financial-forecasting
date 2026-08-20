"""Shared Word-document builder.

Both project documents are generated rather than hand-edited, for the same reason
the results tables are: a document that is regenerated cannot drift away from the
code it describes. This module holds the styling and the tiny markup layer they
share; the documents themselves are content, not formatting.

Supports **bold** and `code` inline markup so the content reads as content.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "BTP_Engineering_Report.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
NAVY_MID = RGBColor(0x2E, 0x5C, 0x8A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x72, 0x80)

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


# --------------------------------------------------------------------- utils
def _shade(paragraph, hex_fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    paragraph._p.get_or_add_pPr().append(shd)


def _border_top(paragraph, hex_color: str = "1F3864", size: int = 8) -> None:
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), "2")
    top.set(qn("w:color"), hex_color)
    pbdr.append(top)
    paragraph._p.get_or_add_pPr().append(pbdr)


class Doc:
    """Thin builder over python-docx so the content below reads as content."""

    def __init__(self) -> None:
        self.d = Document()
        style = self.d.styles["Normal"]
        style.font.name = BODY_FONT
        style.font.size = Pt(10.5)
        style.font.color.rgb = INK
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        for section in self.d.sections:
            section.left_margin = section.right_margin = Inches(0.85)
            section.top_margin = section.bottom_margin = Inches(0.8)

    # -- structure --
    def title(self, text: str, subtitle: str = "", meta: str = "") -> None:
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size = Pt(24)
        r.font.bold = True
        r.font.color.rgb = NAVY
        if subtitle:
            p2 = self.d.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(subtitle)
            r2.font.size = Pt(13)
            r2.font.color.rgb = NAVY_MID
            r2.font.italic = True
        if meta:
            p3 = self.d.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r3 = p3.add_run(meta)
            r3.font.size = Pt(9.5)
            r3.font.color.rgb = MUTED

    def h1(self, text: str) -> None:
        self.d.add_paragraph()
        p = self.d.add_paragraph()
        _border_top(p)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(text)
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def h2(self, text: str) -> None:
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = NAVY_MID

    def h3(self, text: str) -> None:
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = INK

    # -- content --
    def p(self, text: str) -> None:
        """Paragraph with **bold** and `code` inline markup."""
        para = self.d.add_paragraph()
        self._rich(para, text)

    def _rich(self, para, text: str) -> None:
        import re

        for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                para.add_run(token[2:-2]).bold = True
            elif token.startswith("`") and token.endswith("`"):
                r = para.add_run(token[1:-1])
                r.font.name = MONO_FONT
                r.font.size = Pt(9.5)
                r.font.color.rgb = NAVY_MID
            else:
                para.add_run(token)

    def bullets(self, items: list[str], numbered: bool = False) -> None:
        for item in items:
            para = self.d.add_paragraph(style="List Number" if numbered else "List Bullet")
            para.paragraph_format.space_after = Pt(3)
            self._rich(para, item)

    def code(self, text: str) -> None:
        para = self.d.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.2)
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(8)
        _shade(para, "F4F5F7")
        r = para.add_run(text)
        r.font.name = MONO_FONT
        r.font.size = Pt(9)

    def callout(self, text: str) -> None:
        para = self.d.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.15)
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(8)
        _shade(para, "E8F0FA")
        self._rich(para, text)

    def table(self, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None) -> None:
        t = self.d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for i, head in enumerate(headers):
            hdr[i].text = ""
            para = hdr[i].paragraphs[0]
            r = para.add_run(head)
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(para, "1F3864")
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "1F3864")
            hdr[i]._tc.get_or_add_tcPr().append(shd)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(2)
                self._rich(para, str(value))
                for run in para.runs:
                    run.font.size = Pt(9.5)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        self.d.add_paragraph()

    def page_break(self) -> None:
        self.d.add_page_break()

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(path)
        return path
